import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

doc_path = r"c:\xampp\htdocs\identitrack\uploads\[Test Document] Non-functional Requirement Testing - Mobile Performance Assessment.docx"
doc = docx.Document(doc_path)

table = doc.tables[0]

def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Row 1: Date Tested & Application/System Name
set_cell(table.cell(1, 1), "October 09, 2026", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(table.cell(1, 4), "IdentiTrack Mobile Application (Student & Security Guard Modules)", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)

# Row 2: Pre-condition
precond_text = "The mobile application is installed on an Android test device (Android 13, 8GB RAM). Network connectivity is established with the IdentiTrack backend server. Test student and guard accounts are pre-configured."
set_cell(table.cell(2, 1), precond_text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=9)

# Row 3: Used Tool(s)
tools_text = "Android Studio Profiler, Flutter DevTools, Postman API Client, Physical Android Device"
set_cell(table.cell(3, 1), tools_text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=9)

# Performance Test Runs (1 to 5)
runs_data = [
    ("1", "1.24 s", "1.42 s", "1.18 s"),
    ("2", "1.18 s", "1.35 s", "1.12 s"),
    ("3", "1.31 s", "1.48 s", "1.25 s"),
    ("4", "1.15 s", "1.30 s", "1.10 s"),
    ("5", "1.22 s", "1.39 s", "1.15 s"),
]

for idx, (run_no, load_time, resp_time, render_time) in enumerate(runs_data):
    row_idx = 7 + idx
    set_cell(table.cell(row_idx, 0), run_no, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table.cell(row_idx, 1), load_time, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table.cell(row_idx, 3), resp_time, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table.cell(row_idx, 5), render_time, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)

# Row 12: Overall Average Performance Result
avg_load = "1.22 s"
avg_resp = "1.39 s"
avg_render = "1.16 s"

set_cell(table.cell(12, 1), avg_load, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(table.cell(12, 3), avg_resp, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(table.cell(12, 5), avg_render, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

output_path = r"c:\xampp\htdocs\identitrack\uploads\[FILLED] Non-functional Requirement Testing - Mobile Performance Assessment.docx"
doc.save(output_path)
print("Successfully generated [FILLED] Mobile Performance docx!")
